"""Catálogo local de capacidades de ampliación por modelo de equipo."""

import math
import os
import re
import sqlite3
import sys
from datetime import date, datetime
from zipfile import BadZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.utils import get_column_letter


COLUMNAS = (
    "Marca",
    "Modelo",
    "RAM máxima (GB)",
    "Tipo de RAM",
    "Ranuras de RAM",
    "RAM soldada (GB)",
    "SSD máximo (GB)",
    "Interfaces de almacenamiento",
    "Notas",
    "Fuente",
    "Verificado el",
)


def _ruta_datos() -> str:
    if getattr(sys, "frozen", False):
        raiz = os.path.dirname(sys.executable)
    else:
        raiz = os.path.dirname(os.path.dirname(__file__))
    ruta = os.path.join(raiz, "datos_ampliacion")
    os.makedirs(ruta, exist_ok=True)
    return ruta


def _normalizar(texto: str) -> str:
    return re.sub(r"[^a-z0-9]", "", texto.casefold())


class CatalogoAmpliacion:
    """Consulta SQLite e importa libros Excel con compatibilidad de equipos."""

    def __init__(self):
        self.directorio = _ruta_datos()
        self.ruta_bd = os.path.join(self.directorio, "compatibilidad.db")
        self.ruta_plantilla = os.path.join(self.directorio, "plantilla_ampliacion.xlsx")
        with self._conectar() as conexion:
            conexion.execute(
                """
                CREATE TABLE IF NOT EXISTS equipos (
                    id INTEGER PRIMARY KEY,
                    marca TEXT NOT NULL,
                    modelo TEXT NOT NULL,
                    ram_max_gb REAL,
                    tipo_ram TEXT,
                    ranuras_ram TEXT,
                    ram_soldada_gb REAL,
                    ssd_max_gb REAL,
                    interfaces_almacenamiento TEXT,
                    notas TEXT,
                    fuente TEXT,
                    verificado_el TEXT,
                    UNIQUE(marca COLLATE NOCASE, modelo COLLATE NOCASE)
                )
                """
            )

    def _conectar(self):
        return sqlite3.connect(self.ruta_bd)

    @staticmethod
    def detectar_equipo():
        """Devuelve marca y modelo de Windows cuando WMI está disponible."""
        try:
            import wmi

            sistemas = wmi.WMI().Win32_ComputerSystem()
            if not sistemas:
                return None
            marca = (sistemas[0].Manufacturer or "").strip()
            modelo = (sistemas[0].Model or "").strip()
            if not marca or not modelo:
                return None
            return {"marca": marca, "modelo": modelo}
        except (ImportError, OSError, AttributeError):
            return None

    def buscar(self, texto: str):
        termino = _normalizar(texto)
        if not termino:
            return []
        with self._conectar() as conexion:
            filas = conexion.execute(
                """
                SELECT marca, modelo, ram_max_gb, tipo_ram, ranuras_ram,
                       ram_soldada_gb, ssd_max_gb, interfaces_almacenamiento,
                       notas, fuente, verificado_el
                FROM equipos
                ORDER BY marca, modelo
                """
            ).fetchall()
        return [
            self._como_dict(fila)
            for fila in filas
            if termino in _normalizar(f"{fila[0]} {fila[1]}")
        ]

    def buscar_equipo(self, marca: str, modelo: str):
        marca_normalizada = _normalizar(marca)
        modelo_normalizado = _normalizar(modelo)
        with self._conectar() as conexion:
            filas = conexion.execute(
                """
                SELECT marca, modelo, ram_max_gb, tipo_ram, ranuras_ram,
                       ram_soldada_gb, ssd_max_gb, interfaces_almacenamiento,
                       notas, fuente, verificado_el
                FROM equipos
                """
            ).fetchall()
        for fila in filas:
            if (
                _normalizar(fila[0]) == marca_normalizada
                and _normalizar(fila[1]) == modelo_normalizado
            ):
                return self._como_dict(fila)
        return None

    def importar_excel(self, ruta: str) -> int:
        if not os.path.isfile(ruta):
            raise FileNotFoundError(f"No se encontró el archivo: {ruta}")
        try:
            libro = load_workbook(ruta, read_only=True, data_only=True)
        except (InvalidFileException, BadZipFile) as error:
            raise ValueError(
                "El archivo no es un Excel .xlsx válido."
            ) from error
        try:
            hoja = libro["Catalogo"] if "Catalogo" in libro.sheetnames else libro.active
            filas = hoja.iter_rows(values_only=True)
            encabezados = next(filas, None)
            if not encabezados:
                raise ValueError("El Excel está vacío; usa la plantilla de ampliación.")
            indices = {
                str(valor).strip(): indice
                for indice, valor in enumerate(encabezados)
                if valor is not None
            }
            faltantes = [columna for columna in COLUMNAS if columna not in indices]
            if faltantes:
                raise ValueError(
                    "Faltan columnas requeridas: " + ", ".join(faltantes)
                )

            registros = []
            for numero_fila, fila in enumerate(filas, start=2):
                valores = {
                    columna: fila[indices[columna]]
                    if indices[columna] < len(fila)
                    else None
                    for columna in COLUMNAS
                }
                if not any(
                    valor is not None and str(valor).strip()
                    for valor in valores.values()
                ):
                    continue
                marca = self._texto(valores["Marca"])
                modelo = self._texto(valores["Modelo"])
                if not marca or not modelo:
                    raise ValueError(
                        f"La fila {numero_fila} requiere Marca y Modelo."
                    )
                fuente = self._texto(valores["Fuente"])
                verificado = self._texto(valores["Verificado el"])
                ram_maxima = self._numero(
                    valores["RAM máxima (GB)"], "RAM máxima", numero_fila
                )
                ssd_maximo = self._numero(
                    valores["SSD máximo (GB)"], "SSD máximo", numero_fila
                )
                if ram_maxima is None and ssd_maximo is None:
                    raise ValueError(
                        f"La fila {numero_fila} requiere RAM máxima o SSD máximo."
                    )
                if not fuente or not verificado:
                    raise ValueError(
                        f"La fila {numero_fila} requiere Fuente y Verificado el."
                    )
                registros.append(
                    (
                        marca,
                        modelo,
                        ram_maxima,
                        self._texto(valores["Tipo de RAM"]),
                        self._texto(valores["Ranuras de RAM"]),
                        self._numero(
                            valores["RAM soldada (GB)"], "RAM soldada", numero_fila
                        ),
                        ssd_maximo,
                        self._texto(valores["Interfaces de almacenamiento"]),
                        self._texto(valores["Notas"]),
                        fuente,
                        verificado,
                    )
                )
        finally:
            libro.close()

        if not registros:
            raise ValueError("El Excel no contiene equipos para importar.")

        with self._conectar() as conexion:
            conexion.executemany(
                """
                INSERT INTO equipos (
                    marca, modelo, ram_max_gb, tipo_ram, ranuras_ram,
                    ram_soldada_gb, ssd_max_gb, interfaces_almacenamiento,
                    notas, fuente, verificado_el
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                ON CONFLICT(marca, modelo) DO UPDATE SET
                    ram_max_gb = excluded.ram_max_gb,
                    tipo_ram = excluded.tipo_ram,
                    ranuras_ram = excluded.ranuras_ram,
                    ram_soldada_gb = excluded.ram_soldada_gb,
                    ssd_max_gb = excluded.ssd_max_gb,
                    interfaces_almacenamiento = excluded.interfaces_almacenamiento,
                    notas = excluded.notas,
                    fuente = excluded.fuente,
                    verificado_el = excluded.verificado_el
                """,
                registros,
            )
        return len(registros)

    def crear_plantilla(self, sobrescribir: bool = False) -> str:
        if os.path.exists(self.ruta_plantilla) and not sobrescribir:
            raise FileExistsError(self.ruta_plantilla)
        libro = Workbook()
        instrucciones = libro.active
        instrucciones.title = "Instrucciones"
        instrucciones.append(["Catálogo de ampliación de hardware"])
        instrucciones.append(["Captura un equipo por fila en la hoja Catalogo."])
        instrucciones.append(["Usa capacidades en GB y agrega una fuente verificable."])
        instrucciones.append(["La fecha puede capturarse como AAAA-MM-DD."])
        hoja = libro.create_sheet("Catalogo")
        hoja.append(COLUMNAS)
        hoja.freeze_panes = "A2"
        hoja.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNAS))}1"
        for celda in hoja[1]:
            celda.font = Font(color="FFFFFF", bold=True)
            celda.fill = PatternFill("solid", fgColor="553C9A")
        for indice, encabezado in enumerate(COLUMNAS, start=1):
            hoja.column_dimensions[get_column_letter(indice)].width = max(
                16, min(36, len(encabezado) + 4)
            )
        libro.save(self.ruta_plantilla)
        return self.ruta_plantilla

    @staticmethod
    def _texto(valor) -> str:
        if valor is None:
            return ""
        if isinstance(valor, (date, datetime)):
            return valor.strftime("%Y-%m-%d")
        return str(valor).strip()

    @staticmethod
    def _numero(valor, nombre: str, fila: int):
        if valor is None or str(valor).strip() == "":
            return None
        try:
            numero = float(valor)
        except (TypeError, ValueError) as error:
            raise ValueError(
                f"El campo {nombre} debe ser numérico en la fila {fila}."
            ) from error
        if not math.isfinite(numero) or numero < 0:
            raise ValueError(
                f"El campo {nombre} debe ser un número finito no negativo "
                f"(fila {fila})."
            )
        return numero

    @staticmethod
    def _como_dict(fila):
        campos = (
            "marca",
            "modelo",
            "ram_max_gb",
            "tipo_ram",
            "ranuras_ram",
            "ram_soldada_gb",
            "ssd_max_gb",
            "interfaces_almacenamiento",
            "notas",
            "fuente",
            "verificado_el",
        )
        return dict(zip(campos, fila))
